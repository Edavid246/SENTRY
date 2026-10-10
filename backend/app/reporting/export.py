"""Render a DivisionReport as PDF (reportlab) or DOCX (python-docx).

Both run in-process on the standard fonts: nothing is fetched, nothing leaves the machine. The
marking is placed in the header and footer so it prints on every page.
"""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.reporting.division import DivisionReport

PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def filename(report: DivisionReport, extension: str) -> str:
    return f"{report.division}-status-report-DRAFT.{extension}"


def render_pdf(report: DivisionReport) -> bytes:
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.grey)
    heading = ParagraphStyle("heading", parent=base["Heading2"], fontSize=12.5, spaceBefore=14)
    title = ParagraphStyle("title", parent=base["Title"], alignment=0, fontSize=20, leading=24)

    def p(text: str, style: ParagraphStyle = body) -> Paragraph:
        return Paragraph(escape(text), style)

    def frame(canvas, doc) -> None:  # noqa: ANN001 - reportlab callback
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8.5)
        canvas.drawCentredString(A4[0] / 2, A4[1] - 11 * mm, report.marking)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawCentredString(A4[0] / 2, 9 * mm, f"{report.marking}  |  Page {doc.page}")
        canvas.restoreState()

    story: list = [p(report.title, title), p(report.tagline), Spacer(1, 6)]
    banner = Table(
        [[p(report.banner, ParagraphStyle("b", parent=body, fontName="Helvetica-Bold"))]]
    )
    banner.setStyle(
        TableStyle([("BOX", (0, 0), (-1, -1), 1, colors.black), ("PADDING", (0, 0), (-1, -1), 6)])
    )
    story += [banner, Spacer(1, 8)]
    story.append(p(f"Marking: {report.marking}"))
    story.append(p(f"Generated {report.generated_at} for {report.prepared_for}", small))
    for label, value in report.facts:
        story.append(p(f"{label}: {value}"))
    story.append(p("Summary", heading))
    story += [p(line) for line in report.summary]

    for section in report.sections:
        story.append(p(section.title, heading))
        if not section.rows:
            story.append(p(section.empty_text, small))
            continue
        rows = [[p("Item", small), p("Record", small)]]
        for row in section.rows:
            lead = "NEEDS ATTENTION  " if row.flagged else ""
            text = f"<b>{escape(lead)}{escape(row.label)}</b><br/>{escape(row.detail)}"
            rows.append([Paragraph(text, body), p(row.ref, small)])
        table = Table(rows, colWidths=[None, 24 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ]
            )
        )
        story.append(table)
        if section.marking:
            story.append(p(f"Section marking: {section.marking}", small))

    story.append(p("Sources", heading))
    story.append(p("Records: " + (", ".join(report.refs) or "none")))
    story += [Spacer(1, 10), p(report.notice, small)]

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=20 * mm,
        bottomMargin=18 * mm,
        title=report.title,
        author="Defence Gateway (demo)",
    ).build(story, onFirstPage=frame, onLaterPages=frame)
    return buffer.getvalue()


def render_docx(report: DivisionReport) -> bytes:
    doc = Document()
    doc.core_properties.title = report.title
    doc.core_properties.author = "Defence Gateway (demo)"
    section = doc.sections[0]
    for part in (section.header, section.footer):
        paragraph = part.paragraphs[0]
        paragraph.text = report.marking
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.size = Pt(8.5)

    doc.add_heading(report.title, level=0)
    doc.add_paragraph(report.tagline)
    banner = doc.add_paragraph()
    banner.add_run(report.banner).bold = True
    doc.add_paragraph(f"Marking: {report.marking}")
    meta = doc.add_paragraph(f"Generated {report.generated_at} for {report.prepared_for}")
    meta.runs[0].font.size = Pt(8)
    for label, value in report.facts:
        doc.add_paragraph(f"{label}: {value}")

    doc.add_heading("Summary", level=1)
    for line in report.summary:
        doc.add_paragraph(line)

    for part in report.sections:
        doc.add_heading(part.title, level=1)
        if not part.rows:
            doc.add_paragraph(part.empty_text)
            continue
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Item"
        table.rows[0].cells[1].text = "Record"
        for row in part.rows:
            cells = table.add_row().cells
            lead = cells[0].paragraphs[0]
            if row.flagged:
                lead.add_run("NEEDS ATTENTION  ").bold = True
            lead.add_run(row.label).bold = True
            cells[0].add_paragraph(row.detail)
            cells[1].text = row.ref
        if part.marking:
            note = doc.add_paragraph(f"Section marking: {part.marking}")
            note.runs[0].font.size = Pt(8)

    doc.add_heading("Sources", level=1)
    doc.add_paragraph("Records: " + (", ".join(report.refs) or "none"))
    doc.add_paragraph(report.notice).runs[0].font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
