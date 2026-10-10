"""Division reports (Phase E): the dashboard as a document, with export.

Reports are built from the same audited typed tools as the dashboard, so they hold exactly what
the page holds. Expected markings are hand-authored from the seed: the owner's Poctova report
spans CONFIDENTIAL production runs and a CLIENT-C delivery; the Briech report includes Secret
airframe rows (REC-094..097), which Briech's own lead is not cleared for.
"""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfReader
from test_assistant_endpoints import _audit
from test_auth_endpoints import auth_header

BASE = "/api/v1/reports"
BANNER = "DRAFT FOR HUMAN REVIEW - not an approved document"


def _get(client, username: str, path: str):
    return client.get(f"{BASE}/{path}", headers=auth_header(client, username))


def _report(client, username: str, key: str) -> dict:
    response = _get(client, username, key)
    assert response.status_code == 200, response.text
    return response.json()


def test_requires_login(client) -> None:
    assert client.get(f"{BASE}/poctova").status_code == 401
    assert client.get(f"{BASE}/poctova/export?format=pdf").status_code == 401


def test_report_is_a_draft_with_summary_and_sections(client) -> None:
    report = _report(client, "owner", "poctova")
    assert report["title"] == "Poctova status report"
    assert report["banner"] == BANNER
    assert "fictitious" in report["notice"]
    titles = [s["title"] for s in report["sections"]]
    assert titles[:2] == ["Production runs on QC hold", "All production runs"]
    flagged = sum(s["flagged"] for s in report["sections"])
    assert flagged > 0
    assert report["summary"][0].startswith(f"{flagged} items need attention")


def test_report_holds_exactly_what_the_dashboard_holds(client) -> None:
    division = _get(client, "owner", "poctova")
    page = client.get("/api/v1/divisions/poctova", headers=auth_header(client, "owner")).json()
    refs = {r["ref"] for s in page["sections"] for r in s["rows"]}
    assert division.status_code == 200
    assert set(_report(client, "owner", "poctova")["refs"]) == refs


def test_report_is_a_derived_item_and_inherits_the_marking(client) -> None:
    poctova = _report(client, "owner", "poctova")
    assert poctova["classification"] == "confidential"
    assert poctova["compartments"] == ["CLIENT-C"]
    assert poctova["marking"] == "CONFIDENTIAL (CLIENT-C)"
    briech = _report(client, "owner", "briech")
    assert briech["classification"] == "secret"  # the airframe rows REC-094..097
    # the marking names the level as the UI does (Secret is shown as Government-sensitive)
    assert briech["marking"].startswith("GOVERNMENT-SENSITIVE (")
    assert "SECRET" not in briech["marking"]
    assert {"CLIENT-A", "UAS-OPS"} <= set(briech["compartments"])


def test_a_lower_cleared_caller_gets_a_lower_marked_report(client) -> None:
    lead = _report(client, "briech.lead", "briech")
    assert lead["classification"] == "confidential"
    assert "REC-095" not in lead["refs"]
    assert "REC-095" in _report(client, "owner", "briech")["refs"]


def test_a_hidden_division_report_reads_like_a_missing_one(client) -> None:
    hidden = _get(client, "briech.lead", "poctova")
    unknown = _get(client, "briech.lead", "nonesuch")
    assert hidden.status_code == unknown.status_code == 404
    assert hidden.json() == unknown.json()
    for path in ("poctova/export?format=pdf", "nonesuch/export?format=pdf"):
        assert _get(client, "briech.lead", path).status_code == 404


def test_unsupported_export_format_is_rejected(client) -> None:
    assert _get(client, "owner", "poctova/export?format=exe").status_code == 422
    assert _get(client, "owner", "poctova/export").status_code == 422


def test_pdf_export_carries_banner_and_marking_on_every_page(client) -> None:
    response = _get(client, "owner", "briech/export?format=pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "briech-status-report-DRAFT.pdf" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    pages = PdfReader(io.BytesIO(response.content)).pages
    texts = [page.extract_text() for page in pages]
    marking = _report(client, "owner", "briech")["marking"]
    assert all(marking in text for text in texts)
    assert BANNER in texts[0]
    assert "REC-095" in "".join(texts)


def test_docx_export_carries_banner_and_marking(client) -> None:
    response = _get(client, "owner", "poctova/export?format=docx")
    assert response.status_code == 200
    assert zipfile.is_zipfile(io.BytesIO(response.content))
    document = Document(io.BytesIO(response.content))
    marking = "CONFIDENTIAL (CLIENT-C)"
    section = document.sections[0]
    assert section.header.paragraphs[0].text == marking
    assert section.footer.paragraphs[0].text == marking
    body = "\n".join(p.text for p in document.paragraphs)
    assert BANNER in body
    assert any("REC-086" in cell.text for t in document.tables for c in t.rows for cell in c.cells)


def _exports_after(client, seq: int) -> list[dict]:
    return [e for e in _audit(client) if e["seq"] > seq and e["payload"].get("action") == "export"]


def test_an_export_is_audited_and_a_preview_is_not_an_export(client) -> None:
    tip = max(e["seq"] for e in _audit(client, limit=1))
    _get(client, "owner", "poctova")
    assert _exports_after(client, tip) == []
    _get(client, "owner", "poctova/export?format=pdf")
    (exported,) = _exports_after(client, tip)
    payload = exported["payload"]
    assert (payload["actor"], payload["resource"], payload["decision"]) == (
        "owner",
        "report",
        "allow",
    )
    assert (payload["division"], payload["format"]) == ("poctova", "pdf")
    assert payload["classification"] == "confidential"
    assert payload["compartments"] == ["CLIENT-C"]
    assert "REC-086" in payload["record_ids"]
